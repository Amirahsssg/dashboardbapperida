import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

D = json.load(open("data.json"))
F = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="21518E")
HEAD_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
INPUT = Font(name=F, size=10, color="0000FF")
TITLE = Font(name=F, bold=True, size=14, color="21518E")
NOTE = Font(name=F, italic=True, size=9, color="4A5462")
thin = Side(style="thin", color="DDE3EA")

wb = Workbook()

def numfmt(dec):
    return "0" if dec == 0 else "0." + "0" * dec

def sheet(name, headers, rows, widths, note=None, input_cols=()):
    ws = wb.create_sheet(name)
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = NOTE
        r0 = 3
    for c, hdr in enumerate(headers, 1):
        cell = ws.cell(row=r0, column=c, value=hdr)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r0 + 1 + i, column=c, value=v)
            cell.font = INPUT if (c in input_cols and not (isinstance(v, str) and v.startswith("="))) else BODY
            cell.border = Border(bottom=thin)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[r0].height = 30
    ws.freeze_panes = ws.cell(row=r0 + 1, column=3 if len(headers) > 4 else 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    return ws, r0

# ---------- README ----------
ws = wb.active
ws.title = "README"
lines = [
    ("Data simulasi Dashboard Bapperida Kota Palembang", TITLE),
    ("SEMUA ANGKA DI FILE INI ADALAH DATA CONTOH (SIMULASI), BUKAN DATA RESMI.", Font(name=F, bold=True, size=11, color="8E1C1C")),
    ("Dipakai untuk prototipe dashboard internal. Ganti dengan data resmi dari BPS, perangkat daerah dan dokumen perencanaan yang ditetapkan.", BODY),
    ("", BODY),
    ("Isi sheet", Font(name=F, bold=True, size=11)),
    ("Indikator — 30 indikator IKU/IKD/IKK: realisasi 2021–2025, target RPJMD 2025–2029, proyeksi 2029 dan status (dihitung dengan rumus).", BODY),
    ("Program_RKPD_2026 — program pendukung per indikator, pagu indikatif 2026 (Rp juta), realisasi anggaran 2025 dan catatan Monev.", BODY),
    ("Dokumen — daftar dokumen perencanaan kota, perangkat daerah dan regulasi beserta metadata.", BODY),
    ("Target_Dokumen_PD — target indikator yang tertulis di Renstra/Renja PD dibandingkan dokumen rujukan (bahan pemeriksaan otomatis).", BODY),
    ("Temuan_Pemeriksaan — hasil pemeriksaan otomatis: perbedaan target, nama, satuan, atau indikator yang tidak dimuat.", BODY),
    ("DSSD — data statistik sektoral per urusan 2021–2025 (data dasar Buku Saku).", BODY),
    ("Urusan — pembagian urusan per bidang, perangkat daerah pengampu dan isu strategis.", BODY),
    ("", BODY),
    ("Cara membaca status", Font(name=F, bold=True, size=11)),
    ("Proyeksi 2029 = realisasi 2025 + (realisasi 2025 − realisasi 2021), yaitu kelanjutan tren 4 tahun terakhir. Untuk satuan %, dibatasi 0–100.", BODY),
    ("Arah 'naik': Sesuai jalur bila proyeksi ≥ target; Perlu perhatian bila ≥ 97% target; selain itu Berisiko.", BODY),
    ("Arah 'turun': Sesuai jalur bila proyeksi ≤ target; Perlu perhatian bila ≤ 105% target; selain itu Berisiko.", BODY),
    ("Teks biru = nilai yang diisi (data contoh). Teks hitam = rumus. Nama program mengikuti nomenklatur Kepmendagri 050; nomor dan tanggal penetapan dokumen daerah sengaja dikosongkan.", BODY),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 130

# ---------- Indikator ----------
IND = D["indikator"]
hdr = ["Kode", "Indikator", "Jenis", "Satuan", "Arah", "Urusan", "Bidang", "Perangkat daerah", "Sumber data",
       "Realisasi 2021", "Realisasi 2022", "Realisasi 2023", "Realisasi 2024", "Realisasi 2025",
       "Target 2025", "Target 2026", "Target 2027", "Target 2028", "Target 2029",
       "Proyeksi 2029", "Status", "Alasan (bahasa sederhana)", "Saran AI (contoh)"]
bid = {b["id"]: b["name"] for b in D["bidang"]}
rows = []
r0 = 3
for i, d in enumerate(IND):
    r = r0 + 1 + i
    proj = f"=IF(D{r}=\"%\",MAX(0,MIN(100,N{r}+(N{r}-J{r}))),N{r}+(N{r}-J{r}))"
    status = (f"=IF(E{r}=\"turun\",IF(T{r}<=S{r}+1E-9,\"Sesuai jalur\",IF(T{r}<=S{r}*1.05,\"Perlu perhatian\",\"Berisiko\")),"
              f"IF(T{r}>=S{r}-1E-9,\"Sesuai jalur\",IF(T{r}>=S{r}*0.97,\"Perlu perhatian\",\"Berisiko\")))")
    rows.append([d["kode"], d["nama"], d["jenis"], d["satuan"], d["arah"], d["urusan"], bid[d["bidang"]], d["pd"], d["sumber"],
                 *d["real"], *d["targets"], proj, status, d["alasan"], d["saran"]])
ws, r0 = sheet("Indikator", hdr, rows,
               [9, 42, 7, 14, 8, 24, 30, 30, 30, 11, 11, 11, 11, 11, 10, 10, 10, 10, 10, 12, 15, 60, 80],
               note="Data simulasi. Proyeksi dan Status dihitung dengan rumus (lihat README).",
               input_cols=set(range(10, 20)))
for i, d in enumerate(IND):
    r = r0 + 1 + i
    for c in range(10, 21):
        ws.cell(row=r, column=c).number_format = numfmt(d["dec"])
    for c in (22, 23):
        ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")

# ---------- Program ----------
code = {d["id"]: d["kode"] for d in IND}
rows = [[p["id"].upper(), p["nama"], p["pd"], ", ".join(code[x] for x in p["inds"]), p["pagu"], p["realAnggaran"] / 100, p["kendala"] or "—"]
        for p in D["program"]]
n = len(rows)
ws, r0 = sheet("Program_RKPD_2026", ["ID", "Program", "Perangkat daerah", "Indikator terkait", "Pagu indikatif 2026 (Rp juta)",
                                     "Realisasi anggaran 2025", "Catatan Monev (kendala)"], rows,
               [7, 62, 34, 30, 18, 14, 70], note="Data simulasi. Nama program mengikuti nomenklatur Kepmendagri 050; pagu dan realisasi adalah contoh.",
               input_cols={5, 6})
for i in range(n):
    r = r0 + 1 + i
    ws.cell(row=r, column=5).number_format = "#,##0"
    ws.cell(row=r, column=6).number_format = "0.0%"
    ws.cell(row=r, column=7).alignment = Alignment(wrap_text=True, vertical="top")
tr = r0 + n + 1
ws.cell(row=tr, column=4, value="Total pagu").font = Font(name=F, bold=True, size=10)
ws.cell(row=tr, column=5, value=f"=SUM(E{r0 + 1}:E{r0 + n})").font = Font(name=F, bold=True, size=10)
ws.cell(row=tr, column=5).number_format = "#,##0"
ws.cell(row=tr + 1, column=4, value="Rata-rata realisasi").font = Font(name=F, bold=True, size=10)
ws.cell(row=tr + 1, column=6, value=f"=AVERAGE(F{r0 + 1}:F{r0 + n})").font = Font(name=F, bold=True, size=10)
ws.cell(row=tr + 1, column=6).number_format = "0.0%"

# ---------- Dokumen ----------
rows = [[d["id"], d["type"], d["title"], d["owner"], d["period"], d["nomor"], d["tgl"], d["status"], d["versi"], d["upload"],
         "Ya" if d.get("rujukan") else ("Diperiksa" if d.get("pd") else "—")] for d in D["dokumen"]]
sheet("Dokumen", ["ID", "Jenis", "Judul", "Pemilik", "Tahun / periode", "Nomor penetapan", "Tanggal penetapan", "Status", "Versi",
                  "Diunggah", "Peran dalam pemeriksaan"], rows, [16, 10, 70, 36, 16, 30, 16, 12, 8, 14, 18],
      note="Data simulasi. Nomor dan tanggal penetapan dokumen daerah sengaja dikosongkan; regulasi pusat memakai nomor sebenarnya.")

# ---------- Target dokumen PD ----------
dec = {d["nama"]: d["dec"] for d in IND}
deck = {d["kode"]: d["dec"] for d in IND}
rows = [[t["dokumen"], t["pd"], t["kodeIndikator"], t["indikatorDiDokumen"], t["satuanDiDokumen"], t["jenisTarget"],
         t["nilaiDiDokumen"], t["nilaiRujukan"], t["dokumenRujukan"], t["halaman"], t["hasil"]] for t in D["targetDokumen"]]
ws, r0 = sheet("Target_Dokumen_PD", ["Dokumen", "Perangkat daerah", "Kode indikator", "Indikator di dokumen", "Satuan di dokumen",
                                     "Jenis target", "Nilai di dokumen", "Nilai rujukan", "Dokumen rujukan", "Halaman", "Hasil pemeriksaan"],
               rows, [44, 34, 10, 46, 14, 12, 14, 12, 34, 9, 18],
               note="Data simulasi. Beberapa perbedaan sengaja dibuat untuk menguji fitur pemeriksaan dokumen.",
               input_cols={7, 8})
for i, t in enumerate(D["targetDokumen"]):
    for c in (7, 8):
        ws.cell(row=r0 + 1 + i, column=c).number_format = numfmt(deck[t["kodeIndikator"]])

# ---------- Temuan ----------
rows = [[f["id"].upper(), f["chkDoc"], f["jenis"], f["title"], f["ind"], f["label"], f["refDoc"], f["refPage"], f["refVal"],
         f["chkPage"], f["chkVal"], "Menunggu tinjauan", f["ai"]] for f in D["temuan"]]
ws, r0 = sheet("Temuan_Pemeriksaan", ["ID", "Dokumen diperiksa", "Jenis temuan", "Temuan", "Indikator", "Yang dibandingkan",
                                      "Dokumen rujukan", "Hal. rujukan", "Nilai rujukan", "Hal. dokumen", "Nilai di dokumen",
                                      "Status tinjauan", "Catatan AI (contoh)"], rows,
               [6, 40, 16, 38, 36, 16, 34, 16, 26, 12, 26, 18, 80], note="Data simulasi. Hasil pemeriksaan otomatis saat dokumen diunggah.")
for i in range(len(rows)):
    ws.cell(row=r0 + 1 + i, column=13).alignment = Alignment(wrap_text=True, vertical="top")

# ---------- DSSD ----------
unames = {u["id"]: u["name"] for u in D["urusan"]}
rows = [[x["kode"], unames[x["urusanId"]], x["uraian"], x["satuan"], *x["values"], x["produsen"]] for x in D["dssd"]]
ws, r0 = sheet("DSSD", ["Kode DSSD", "Urusan", "Uraian", "Satuan", "2021", "2022", "2023", "2024", "2025", "Produsen data"], rows,
               [14, 30, 38, 12, 11, 11, 11, 11, 11, 36], note="Data simulasi.", input_cols=set(range(5, 10)))
for i, x in enumerate(D["dssd"]):
    for c in range(5, 10):
        ws.cell(row=r0 + 1 + i, column=c).number_format = "#,##0" if x["dec"] == 0 else "#,##0." + "0" * x["dec"]

# ---------- Urusan ----------
rows = [[u["name"], bid[u["bidang"]], u["pd"], len(u["inds"]), "; ".join(u["isu"])] for u in D["urusan"]]
sheet("Urusan", ["Urusan", "Bidang pengampu", "Perangkat daerah utama", "Jumlah indikator", "Isu strategis (contoh)"], rows,
      [32, 36, 36, 14, 80], note="Pembagian urusan adalah contoh dan perlu disesuaikan dengan pembagian tugas resmi setiap bidang.")

# ---------- Realisasi anggaran RKPD 2026 (bulanan, kumulatif) ----------
A = D["anggaran2026"]
quarter = {"Mar": "Triwulan I", "Jun": "Triwulan II", "Sep": "Triwulan III"}
rows = []
r0 = 3
for i, b in enumerate(A["bulan"]):
    r = r0 + 1 + i
    rows.append([b + " 2026", quarter.get(b, ""), A["rencana"][i] / 100, A["realisasi"][i] / 100,
                 f"=D{r}-C{r}", f"=D{r}*'Program_RKPD_2026'!$E$29"])
ws, r0 = sheet("Realisasi_Anggaran_2026", ["Bulan", "Akhir triwulan", "Rencana kumulatif", "Realisasi kumulatif", "Selisih",
                                           "Realisasi (Rp juta)"], rows, [14, 16, 18, 18, 12, 20],
               note="Data simulasi. Persentase kumulatif terhadap total pagu indikatif program di sheet Program_RKPD_2026.",
               input_cols={3, 4})
for i in range(len(rows)):
    r = r0 + 1 + i
    for c in (3, 4, 5):
        ws.cell(row=r, column=c).number_format = "0.0%"
    ws.cell(row=r, column=6).number_format = "#,##0"

wb.save("Data_Simulasi_Dashboard_Bapperida.xlsx")
print("saved")
